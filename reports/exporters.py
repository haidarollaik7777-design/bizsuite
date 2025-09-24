from io import BytesIO
from decimal import Decimal
from django.http import HttpResponse
from django.apps import apps
from django.conf import settings
from django.db.models import Sum, Value as V, DecimalField
from django.db.models.functions import Coalesce
from openpyxl import Workbook

def _resolve_model_and_fields():
    Model = None
    account_field, debit_field, credit_field, date_lookup = "account","debit","credit","date"
    try:
        Model = apps.get_model("ledger","JournalLine")
    except Exception:
        Model = None
    if Model is None:
        conf = getattr(settings, "REPORTS_POSTING", {}) or {}
        mpath = conf.get("model","")
        if "." in mpath:
            app_label, model_name = mpath.split(".")
            try:
                Model = apps.get_model(app_label, model_name)
            except Exception:
                Model = None
        account_field = conf.get("account", account_field)
        debit_field   = conf.get("debit",  debit_field)
        credit_field  = conf.get("credit", credit_field)
        date_lookup   = conf.get("date",   date_lookup)
    if Model is None:
        return None, account_field, debit_field, credit_field, None
    try:
        fnames = {f.name for f in Model._meta.get_fields()}
    except Exception:
        fnames = set()
    if "date" not in fnames and "entry" in fnames:
        date_lookup = "entry__date"
    elif "date" not in fnames:
        date_lookup = None
    return Model, account_field, debit_field, credit_field, date_lookup

def _apply_filters(request, qs, date_lookup, account_field):
    df = request.GET.get("date_from"); dt = request.GET.get("date_to")
    acc_code = request.GET.get("account_code") or request.GET.get("code") or ""
    acc_id   = request.GET.get("account")
    try:
        if date_lookup and df: qs = qs.filter(**{f"{date_lookup}__gte": df})
        if date_lookup and dt: qs = qs.filter(**{f"{date_lookup}__lte": dt})
        if acc_code: qs = qs.filter(**{f"{account_field}__code__startswith": acc_code})
        if acc_id:   qs = qs.filter(**{f"{account_field}__id": acc_id})
    except Exception:
        pass
    return qs

def _wb_response(wb, filename):
    bio = BytesIO(); wb.save(bio); bio.seek(0)
    resp = HttpResponse(bio.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp

def gl_xlsx(request):
    Model, account_field, debit_field, credit_field, date_lookup = _resolve_model_and_fields()
    if Model is None:
        return HttpResponse("Posting model not found", status=500)
    try:
        qs = Model.objects.select_related(account_field, "entry")
    except Exception:
        qs = Model.objects.all()
    qs = _apply_filters(request, qs, date_lookup, account_field)
    for ob in (["date","id"], [f"{date_lookup}","id"] if date_lookup else [], ["id"]):
        try:
            if ob: qs = qs.order_by(*ob); break
        except Exception:
            continue
    wb = Workbook(); ws = wb.active; ws.title = "General Ledger"
    ws.append(["Date","Code","Name","Debit","Credit","Memo"])
    for obj in qs:
        acct = getattr(obj, account_field, None)
        code = getattr(acct, "code", ""); name = getattr(acct, "name", "")
        memo = getattr(obj, "memo", "") or getattr(obj, "label", "") or ""
        dval = getattr(obj, "date", None)
        if not dval:
            ent = getattr(obj, "entry", None); dval = getattr(ent, "date", "")
        dr = getattr(obj, debit_field, Decimal("0")) or Decimal("0")
        cr = getattr(obj, credit_field, Decimal("0")) or Decimal("0")
        ws.append([dval, code, name, float(dr), float(cr), memo])
    return _wb_response(wb, "general_ledger.xlsx")

def tb_xlsx(request):
    Model, account_field, debit_field, credit_field, date_lookup = _resolve_model_and_fields()
    if Model is None:
        return HttpResponse("Posting model not found", status=500)
    qs = _apply_filters(request, Model.objects.all(), date_lookup, account_field)
    code_path = f"{account_field}__code"; name_path = f"{account_field}__name"
    ann = qs.values(code_path, name_path).annotate(
        dr=Coalesce(Sum(debit_field), V(Decimal("0.00")), output_field=DecimalField(max_digits=18, decimal_places=2)),
        cr=Coalesce(Sum(credit_field),V(Decimal("0.00")), output_field=DecimalField(max_digits=18, decimal_places=2)),
    ).order_by(code_path)
    wb = Workbook(); ws = wb.active; ws.title="Trial Balance"
    ws.append(["Code","Name","Debit","Credit","Balance"])
    tot_dr = tot_cr = Decimal("0")
    for a in ann:
        code = str(a.get(code_path) or ""); name = a.get(name_path) or ""
        if not code: continue
        bal = (a["dr"] or Decimal("0")) - (a["cr"] or Decimal("0"))
        debit_col  = bal if bal > 0 else Decimal("0")
        credit_col = -bal if bal < 0 else Decimal("0")
        tot_dr += debit_col; tot_cr += credit_col
        ws.append([code, name, float(debit_col), float(credit_col), float(bal)])
    ws.append(["","TOTALS", float(tot_dr), float(tot_cr), float(tot_dr - tot_cr)])
    return _wb_response(wb, "trial_balance.xlsx")

def bs_xlsx(request):
    Model, account_field, debit_field, credit_field, date_lookup = _resolve_model_and_fields()
    if Model is None:
        return HttpResponse("Posting model not found", status=500)
    qs = Model.objects.all()
    dt = request.GET.get("date_to")
    if date_lookup and dt:
        try: qs = qs.filter(**{f"{date_lookup}__lte": dt})
        except Exception: pass
    code_path = f"{account_field}__code"; name_path = f"{account_field}__name"
    ann = qs.values(code_path, name_path).annotate(
        dr=Coalesce(Sum(debit_field), V(Decimal("0.00")), output_field=DecimalField(max_digits=18, decimal_places=2)),
        cr=Coalesce(Sum(credit_field),V(Decimal("0.00")), output_field=DecimalField(max_digits=18, decimal_places=2)),
    ).order_by(code_path)
    assets=[]; liab=[]; equity=[]; totA=totL=totE=Decimal("0")
    rev = exp = Decimal("0")
    for a in ann:
        code = str(a.get(code_path) or "")
        if not code: continue
        name = a.get(name_path) or ""
        head = code[0]
        dr = a["dr"] or Decimal("0"); cr = a["cr"] or Decimal("0")
        if head=="1":
            amt = dr - cr
            if amt: assets.append([code,name,float(amt)]); totA += amt
        elif head=="2":
            amt = cr - dr
            if amt: liab.append([code,name,float(amt)]);  totL += amt
        elif head=="3":
            amt = cr - dr
            if amt: equity.append([code,name,float(amt)]); totE += amt
        if head in ("4","7"): rev += (cr - dr)
        elif head in ("5","6","8","9"): exp += (dr - cr)
    re = rev - exp
    if re: equity.append(["", "Retained Earnings (to date)", float(re)]); totE += re
    wb = Workbook(); ws = wb.active; ws.title = "Balance Sheet"
    ws.append(["Section","Code","Name","Amount"])
    for r in assets: ws.append(["Assets"]+r)
    ws.append(["Totals","","Total Assets", float(totA)])
    for r in liab:   ws.append(["Liabilities"]+r)
    for r in equity: ws.append(["Equity"]+r)
    ws.append(["Totals","","Total Liabilities + Equity", float(totL + totE)])
    return _wb_response(wb, "balance_sheet.xlsx")

def is_xlsx(request):
    Model, account_field, debit_field, credit_field, date_lookup = _resolve_model_and_fields()
    if Model is None:
        return HttpResponse("Posting model not found", status=500)
    qs = _apply_filters(request, Model.objects.all(), date_lookup, account_field)
    code_path = f"{account_field}__code"; name_path = f"{account_field}__name"
    ann = qs.values(code_path, name_path).annotate(
        dr=Coalesce(Sum(debit_field), V(Decimal("0.00")), output_field=DecimalField(max_digits=18, decimal_places=2)),
        cr=Coalesce(Sum(credit_field),V(Decimal("0.00")), output_field=DecimalField(max_digits=18, decimal_places=2)),
    ).order_by(code_path)
    wb = Workbook(); ws = wb.active; ws.title="Income Statement"
    ws.append(["Section","Code","Name","Amount"])
    tot_rev = tot_exp = Decimal("0")
    for a in ann:
        code = str(a.get(code_path) or "")
        if not code: continue
        name = a.get(name_path) or ""
        head = code[0]
        dr = a["dr"] or Decimal("0"); cr = a["cr"] or Decimal("0")
        if head in ("4","7"):
            amt = cr - dr
            if amt: ws.append(["Revenue", code, name, float(amt)]); tot_rev += amt
        elif head in ("5","6","8","9"):
            amt = dr - cr
            if amt: ws.append(["Expenses", code, name, float(amt)]); tot_exp += amt
    net = tot_rev - tot_exp
    ws.append(["Totals","","Total Revenue", float(tot_rev)])
    ws.append(["Totals","","Total Expenses", float(tot_exp)])
    ws.append(["Totals","","Net Income", float(net)])
    return _wb_response(wb, "income_statement.xlsx")
