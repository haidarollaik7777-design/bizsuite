from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def rows_to_xlsx_response(filename: str, headers: list[str], rows: list[list]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for col_idx, _ in enumerate(headers, start=1):
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            val = row[0].value
            if val is None:
                continue
            max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    safe = filename.replace(" ", "_").replace("/", "_")
    resp["Content-Disposition"] = f'attachment; filename="{safe}.xlsx"'
    wb.save(resp)
    return resp
